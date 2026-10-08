"""泰國工業經濟辦公室（OIE，工業部）工業指數月報 xlsx：製造業生產指數 MPI、產能利用率、出貨、庫存（免金鑰）。

檔案網址：https://www.oie.go.th/assets/portals/1/fileups/2/files/Industrial%20index/indexes_eng/month/<file>.xlsx
版面（第一張工作表）：年份列（'2021'、之後每年第一個月欄標示）＋下一列月份（'JAN.'…'DEC.'，最新月可能寫成 'AUG.*'）；
列名在前三欄之一；最後兩欄是 % Change，不是期別。所有指數基期 2021=100，檔內只有 2021 年 1 月起的資料。
同一個檔只下載一次、多條序列共用；同站請求至少隔 1 秒。

params：
  file       檔名（不含 .xlsx），如 "Prodidx1_E"（生產）、"capidx_E"（產能利用率）、"Shipidx_E"、"Invidx_E"（必填）
  row_label  列名開頭（必填），如 "Integrated Index (Seasonally adjusted)"、"TSIC : 26 Manufacture of computer, electronic and optical products"
  scale      選填
  probe_url  選填（fetcher 不讀）
"""
from __future__ import annotations

import io
import re
import time

try:
    from . import _http
    from . import an_common as A
except ImportError:
    import _http
    import an_common as A

BASE = "https://www.oie.go.th/assets/portals/1/fileups/2/files/Industrial%20index/indexes_eng/month/"
MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def _month_of(cell):
    if cell is None:
        return None
    t = str(cell).strip().upper().rstrip("*").rstrip(".").strip()
    return MON.index(t) + 1 if t in MON else None


def _year_of(cell):
    if isinstance(cell, (int, float)) and 1900 < cell < 2200:
        return int(cell)
    if isinstance(cell, str) and cell.strip().isdigit() and len(cell.strip()) == 4:
        return int(cell.strip())
    return None


def parse_columns(rows):
    """回傳 (資料起始列索引, [(欄索引, 'YYYY-MM-01')])。找不到月份列就丟 ValueError。"""
    mi = next((i for i, r in enumerate(rows[:12]) if sum(1 for c in r if _month_of(c)) >= 6), None)
    if mi is None or mi == 0:
        raise ValueError("找不到月份表頭列（版面改版？）")
    cols, y = [], None
    for j, (a, b) in enumerate(zip(rows[mi - 1], rows[mi])):
        ya = _year_of(a)
        if ya:
            y = ya
        m = _month_of(b)
        if m and y:
            cols.append((j, A.ym_date(y, m)))
    return mi + 1, cols


def row_name(r):
    names = [str(x).strip() for x in r[:3] if isinstance(x, str) and x.strip() and not re.match(r"^[\d.]+$", x.strip())]
    return re.sub(r"\s+", " ", names[-1]) if names else ""


def extract(rows, start, cols, label, scale=1.0):
    want = re.sub(r"\s+", " ", label).strip().lower()
    for r in rows[start:]:
        if row_name(r).lower().startswith(want):
            return A.clean_obs([(d, r[j]) for j, d in cols if j < len(r) and r[j] not in (None, "")], scale)
    raise ValueError("找不到列「%s」" % label)


def load_rows(content: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def fetch(specs: list[dict]) -> dict:
    by_file: dict = {}
    for s in specs:
        by_file.setdefault(s["params"]["file"], []).append(s)
    res: dict = {}
    for i, (f, group) in enumerate(by_file.items()):
        if i:
            time.sleep(1.1)
        try:
            rows = load_rows(_http.get("%s%s.xlsx" % (BASE, f), binary=True))
            start, cols = parse_columns(rows)
        except Exception as e:  # noqa: BLE001
            for s in group:
                res[s["sid"]] = {"error": "OIE %s：%s" % (f, str(e)[:160])}
            continue
        for s in group:
            p = s["params"]
            try:
                obs = extract(rows, start, cols, p["row_label"], float(p.get("scale", 1.0)))
                res[s["sid"]] = {"obs": obs} if obs else {"error": "解析後沒有任何觀測值"}
            except Exception as e:  # noqa: BLE001
                res[s["sid"]] = {"error": "OIE %s：%s" % (f, str(e)[:160])}
    return res
