"""馬來西亞國家銀行（BNM）月度統計（Monthly Highlights & Statistics）xlsx。

BNM 每月發布新一期時，xlsx 會放進新的資料夾（2026-07 是 22576999，2026-09 起是 23200938），舊資料夾不再更新，
所以不能寫死網址。做法：
  1. 抓 BNM 首頁，找出最新一期的頁面（/-/monthly-highlights-statistics-in-<月>-<年>）；
  2. 讀該頁的表格清單，取得各表號（如 2.5）目前的 xlsx 連結；
  3. 解析不到時，退回 params.url（固定網址，可能已不再更新，頁面會標過期）。
params：
  table   表號，如 "2.5"（必填）
  layout  "ym"（A 欄年、B 欄月 1–12）／"yq"（A 欄年、B 欄 1Q–4Q）／"daily"（B 欄 'June 2026'、C 欄日）
  col     取值欄位（0 起算）
  url     備援固定網址（也供 CI --probe 使用）
  scale   選填
同一張表只下載一次，多條序列共用；請求之間隔 1 秒。
"""
from __future__ import annotations

import datetime
import io
import re
import time

try:
    from . import _http
    from . import an_common as A
except ImportError:
    import _http
    import an_common as A

HOME = "https://www.bnm.gov.my/"
BASE = "https://www.bnm.gov.my"
EDITION_RX = re.compile(r'(?:https://www\.bnm\.gov\.my)?(/-/monthly-highlights-(?:and-)?statistics-in-[a-z0-9\-]+)')
LINK_RX = re.compile(r'href="(/documents/20124/\d+/([\w.]+)\.xlsx)"')
GAP = 1.0


def find_edition(home_html: str) -> str | None:
    """首頁上第一個『Monthly Highlights & Statistics』頁面路徑（首頁依新到舊排）。"""
    m = EDITION_RX.search(home_html)
    return m.group(1) if m else None


def table_links(edition_html: str) -> dict:
    """{表號: 絕對網址}；同表號多個連結取第一個。"""
    out: dict = {}
    for path, table in LINK_RX.findall(edition_html):
        out.setdefault(table, BASE + path)
    return out


def discover(get=None) -> dict:
    """回傳 {表號: 網址}；任一步失敗回 {}。"""
    get = get or (lambda u: _http.get(u))
    try:
        ed = find_edition(get(HOME))
        if not ed:
            return {}
        time.sleep(GAP)
        return table_links(get(BASE + ed))
    except Exception:  # noqa: BLE001
        return {}


def _num(x):
    return A.to_float(x)


def parse_ym(rows, col):
    out, yr = [], None
    for r in rows:
        a, b = r[0], r[1]
        if a is not None and re.fullmatch(r"\d{4}", str(a).strip()):
            yr = int(str(a).strip())
        if yr and b is not None and re.fullmatch(r"\d{1,2}", str(b).strip()):
            v = _num(r[col]) if col < len(r) else None
            if v is not None:
                out.append((A.ym_date(yr, int(str(b).strip())), v))
    return out


def parse_yq(rows, col):
    out, yr = [], None
    for r in rows:
        a, b = r[0], r[1]
        if a is not None:
            m = re.match(r"(\d{4})", str(a).strip())
            if m:
                yr = int(m.group(1))
        if yr and b is not None:
            m = re.fullmatch(r"([1-4])Q", str(b).strip())
            if m:
                v = _num(r[col]) if col < len(r) else None
                if v is not None:
                    out.append((A.yq_date(yr, int(m.group(1))), v))
    return out


def parse_daily(rows, col):
    out, ym = [], None
    for r in rows:
        b, c = r[1], r[2]
        if b is not None:
            try:
                ym = datetime.datetime.strptime(str(b).strip(), "%B %Y")
            except ValueError:
                pass
        if ym and c is not None and re.fullmatch(r"\d{1,2}", str(c).strip()):
            v = _num(r[col]) if col < len(r) else None
            if v is not None:
                out.append(("%04d-%02d-%02d" % (ym.year, ym.month, int(str(c).strip())), v))
    return out


PARSERS = {"ym": parse_ym, "yq": parse_yq, "daily": parse_daily}


def sheet_rows(content: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return list(wb[wb.sheetnames[0]].iter_rows(values_only=True))


def fetch(specs: list[dict]) -> dict:
    links = discover()
    rows_by_table: dict = {}
    for s in specs:  # 依序下載每張表一次
        p = s["params"]
        t = p["table"]
        if t in rows_by_table:
            continue
        urls = [u for u in (links.get(t), p.get("url")) if u]
        err = None
        for u in dict.fromkeys(urls):
            try:
                time.sleep(GAP)
                rows_by_table[t] = sheet_rows(_http.get(u, binary=True))
                break
            except Exception as e:  # noqa: BLE001
                err = e
        else:
            rows_by_table[t] = err or RuntimeError("沒有可用網址")

    def one(s):
        p = s["params"]
        rows = rows_by_table[p["table"]]
        if isinstance(rows, Exception):
            raise rows
        raw = PARSERS[p["layout"]](rows, int(p["col"]))
        sc = float(p.get("scale", 1.0))
        return A.clean_obs(raw, sc)

    return A.run_specs(specs, one)
