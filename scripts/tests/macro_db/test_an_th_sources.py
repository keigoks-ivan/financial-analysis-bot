"""泰國 fetcher（an_th_bot／an_th_tpso／an_th_oie）離線解析測試。"""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import pytest

from macro_db.sources import an_th_bot, an_th_oie, an_th_tpso

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "macro_db" / "an" / "th"


def fx(name):
    return (FIX / name).read_text(encoding="utf-8")


# ---------------- BOT
def test_bot_period_date():
    assert an_th_bot.period_date("AUG 2026 p") == "2026-08-01"
    assert an_th_bot.period_date("JUL 2026 e") == "2026-07-01"
    assert an_th_bot.period_date("Q2/2026") is None and an_th_bot.period_date("") is None


def test_bot_parse_and_extract():
    rows, hi, meta = an_th_bot.parse_csv(fx("bot_991.csv"))
    assert meta["updated"].startswith("30 Sep 2026")
    obs = an_th_bot.extract(rows, hi, 1, "Retail Sales Index")
    assert obs[0] == ("2026-02-01", 182.59) and obs[-1] == ("2026-07-01", 129.92)
    # 列名前面的縮排空白不影響比對
    assert an_th_bot.extract(rows, hi, 2, "Non-durable Goods")[-1][1] == 141.48
    assert an_th_bot.extract(rows, hi, 8, "Durable Goods", 2.0)[-1][1] == 232.04


def test_bot_label_mismatch_and_missing_row_are_errors():
    rows, hi, _ = an_th_bot.parse_csv(fx("bot_991.csv"))
    with pytest.raises(ValueError, match="不符"):
        an_th_bot.extract(rows, hi, 2, "Durable Goods")
    with pytest.raises(ValueError, match="找不到列號"):
        an_th_bot.extract(rows, hi, 99, "x")
    with pytest.raises(ValueError):
        an_th_bot.parse_csv("<html>not csv</html>")


def test_bot_skips_missing_values():
    text = ("a\nb\nc\nLast Updated : 1 Jan 2026\nd\n,,AUG 2026,JUL 2026,JUN 2026\n"
            "1,Foo,n.a.,....,3.5\n")
    rows, hi, _ = an_th_bot.parse_csv(text)
    assert an_th_bot.extract(rows, hi, 1, "foo") == [("2026-06-01", 3.5)]


def test_bot_form_range():
    h = fx("bot_form.html")
    d = an_th_bot._form_range(h)
    assert d["__VIEWSTATE"] == "abc&def"            # HTML 跳脫要還原
    assert d["drpFromYear"] == "2011xxxx" and d["drpToYear"] == "2026xxxx"   # 最早年到最新年
    assert d["drpFromMonth"] == "xxxx01xx" and d["drpToMonth"] == "xxxx08xx" and d["drpPeriod"] == "MTH"
    with pytest.raises(ValueError):
        an_th_bot._form_range("<html>no form</html>")


def test_bot_fetch_shares_one_export_and_reports_errors(monkeypatch):
    calls = []

    def fake_export(sess, rid):
        calls.append(rid)
        if rid == "000":
            raise ValueError("找不到起迄年下拉選單")
        return fx("bot_991.csv")

    monkeypatch.setattr(an_th_bot, "export_report", fake_export)
    specs = [{"sid": "a", "params": {"report_id": 991, "row_no": 1, "row_label": "Retail Sales Index"}},
             {"sid": "b", "params": {"report_id": 991, "row_no": 2, "row_label": "Non-durable Goods"}},
             {"sid": "c", "params": {"report_id": 991, "row_no": 2, "row_label": "WRONG"}},
             {"sid": "d", "params": {"report_id": "000", "row_no": 1, "row_label": "x"}}]
    res = an_th_bot.fetch(specs)
    assert calls == ["991", "000"]                       # 同一張報表只匯出一次
    assert res["a"]["obs"][-1] == ("2026-07-01", 129.92) and res["b"]["obs"]
    assert "不符" in res["c"]["error"] and "下拉選單" in res["d"]["error"]


# ---------------- TPSO
def test_tpso_month_parse_be_year():
    out = an_th_tpso.parse_month_rows(json.loads(fx("tpso_month.json")))
    assert out["00000"] == [("2026-09-01", 102.93)]          # 佛曆 2569 -> 西元 2026
    assert "92000" not in out                                  # index 為 null 的略過


def test_tpso_range_parse_and_usd_field():
    res = json.loads(fx("tpso_range.json"))
    assert an_th_tpso.parse_range(res)["00000"][0] == ("1987-01-01", 34.61)
    assert an_th_tpso.parse_range(res, "indexD")["000"] == [("2026-08-01", 114.4)]


def test_tpso_cci_parse_national_rows_only():
    rows = json.loads(fx("tpso_cci.json"))
    for r in rows:
        r["year"] = 2569
    obs = an_th_tpso.parse_cci(rows, "รวม")
    assert obs == [("2026-01-01", 50.1), ("2026-02-01", 49.5), ("2026-09-01", 48.9)]   # 'februay' 是來源的拼字
    assert an_th_tpso.parse_cci(rows, "ปัจจุบัน") == [("2026-01-01", 55.0)]


def test_tpso_recent_months():
    assert an_th_tpso.recent_months(dt.date(2026, 10, 8), 4) == [(2569, 7), (2569, 8), (2569, 9), (2569, 10)]
    assert an_th_tpso.recent_months(dt.date(2026, 1, 15), 3) == [(2568, 11), (2568, 12), (2569, 1)]


def test_tpso_fetch_groups_requests(monkeypatch):
    posted = []

    def fake_post(sess, url, body, tries=3):
        posted.append((url, body.get("type"), tuple(body["commodities"])))
        return json.loads(fx("tpso_month.json"))

    monkeypatch.delenv("MACRO_DB_BACKFILL", raising=False)
    monkeypatch.setattr(an_th_tpso, "_post", fake_post)
    specs = [{"sid": "a", "params": {"api": "cpig", "type": "TG", "code": "00000", "year_base": 2566}},
             {"sid": "b", "params": {"api": "cpig", "type": "TG", "code": "91000", "year_base": 2566}},
             {"sid": "c", "params": {"api": "cpig", "type": "TG", "code": "92000", "year_base": 2566}}]
    res = an_th_tpso.fetch(specs)
    assert len(posted) == an_th_tpso.RECENT_MONTHS            # 三條序列共用每個月一次請求
    assert all(p[0].endswith("/OpenApi/Cpig/Month") and p[2] == ("00000", "91000", "92000") for p in posted)
    assert res["a"]["obs"] == [("2026-09-01", 102.93)] and res["b"]["obs"][0][1] == 101.77
    assert "error" in res["c"]


# ---------------- OIE
def _rows():
    import openpyxl
    wb = openpyxl.load_workbook(FIX / "oie_prod.xlsx", read_only=True, data_only=True)
    return list(wb[wb.sheetnames[0]].iter_rows(values_only=True))


def test_oie_columns_skip_percent_change_and_star():
    rows = _rows()
    start, cols = an_th_oie.parse_columns(rows)
    assert [d for _, d in cols] == ["2021-01-01", "2021-02-01", "2021-03-01", "2022-01-01", "2022-02-01", "2022-08-01"]
    obs = an_th_oie.extract(rows, start, cols, "Integrated Index (Seasonally adjusted)")
    assert obs[0] == ("2021-01-01", 97.5) and obs[-1] == ("2022-08-01", 102.5)     # 'AUG.*' 的星號要去掉
    food = an_th_oie.extract(rows, start, cols, "TSIC : 10 Manufacture of food products")
    assert len(food) == 5                                                          # 空白格略過


def test_oie_missing_row_is_error():
    rows = _rows()
    start, cols = an_th_oie.parse_columns(rows)
    with pytest.raises(ValueError):
        an_th_oie.extract(rows, start, cols, "No such row")
    with pytest.raises(ValueError):
        an_th_oie.parse_columns([[None] * 5] * 6)


def test_oie_fetch_downloads_each_file_once(monkeypatch):
    got = []
    data = (FIX / "oie_prod.xlsx").read_bytes()
    monkeypatch.setattr(an_th_oie._http, "get", lambda url, **kw: got.append(url) or data)
    specs = [{"sid": "a", "params": {"file": "Prodidx1_E", "row_label": "Integrated Index (Seasonally adjusted)"}},
             {"sid": "b", "params": {"file": "Prodidx1_E", "row_label": "Integrated Index (Not seasonally adjusted)"}},
             {"sid": "c", "params": {"file": "Prodidx1_E", "row_label": "nope"}}]
    res = an_th_oie.fetch(specs)
    assert len(got) == 1 and got[0].endswith("Prodidx1_E.xlsx")
    assert res["a"]["obs"] and res["b"]["obs"] and "error" in res["c"]
